"""pptx-ai-slide-forge :: reference deck builder (python-pptx).

This exact script built a real 8-slide deck (NovaPay Liquid Electric Blue)
and passed the full QA loop: geometry lint, markitdown text checks, Keynote
render inspection. Copy it per project and adapt - do not rewrite.

HOW TO ADAPT
1. Theme tokens are the constant block right below this docstring. Replace
   the HEX values with your style spec (references/styles.md); change FONT
   only if your style demands it ("Helvetica Neue" is the safe default).
2. Helper library (txt, card, accent_card, pill, dot, rect, footer, verify)
   encodes QA-relevant invariants: glass alpha via DrawingML XML, letter
   spacing via the `spc` attribute, table cell borders in schema order,
   geometry lint with a 0.05 sq-in overlap threshold. Keep them unchanged.
3. Canvas: 13.333 x 7.5 in = 1920x1080 px @144 dpi. 1 px at 1920 = 0.5 pt.
4. Slide sections are marked with banners - replace their content with your
   outline (action titles!). Backgrounds come from assets/bg_NN.png generated
   by make_textures.py.

Deck content below is the worked example (Russian, "Spec-Driven Development
with AI agents") - keep it as a pattern reference for layout coordinates.
"""
import os

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(BASE, "assets")
OUT_PPTX = os.path.join(BASE, "spec-driven-development-ai-agents.pptx")


def _require_assets():
    if not os.path.isdir(ASSETS):
        raise SystemExit(
            f"backgrounds not found: {ASSETS}\n"
            "generate them first: python3 ../../scripts/make_textures.py "
            "--config <config.py> --out ../../assets")


EMU_IN = 914400


def IN(v):
    return Emu(int(round(v * EMU_IN)))


# ---- NovaPay Liquid Electric Blue tokens ----
BG = "03091B"          # deep black-blue
BG2 = "0A1830"         # chip fill base
BG3 = "0F1C3E"         # table header
GLASS = "20365B"       # smoked blue glass fill (alpha applied)
GLASS_LINE = "8AB4FF"  # luminous glass border (alpha applied)
CARD_ALT = "142A4D"    # opaque dark-navy panel (table rows)
BORDER = "2E4370"      # structural borders / dividers
DIVIDER = "22335A"
BLUE = "075BFF"        # electric blue
BLUE_BRIGHT = "287BFF" # luminous blue (glows)
ULTRA = "2438F5"       # highlight fill
ULTRA_LINE = "C6D6FF"  # highlight border
ICE = "B5D7FF"
SKY = "5BA7FF"         # light data accent
WHITE = "F8FAFF"
COOL = "A5B7D9"        # secondary text
MUTED = "6F82A8"       # muted text / footers
BODY = "E2E9FA"        # body text
BTN_FILL = "0F1C3E"
BTN_BORDER = "5F82C0"

FONT = "Helvetica Neue"
RADIUS = 0.125  # 18 px at 144 px/in

TITLE_W = 12.233
ML = 0.55


def RGB(h):
    return RGBColor.from_string(h)


def set_spc(run, spc):
    run._r.get_or_add_rPr().set("spc", str(int(spc)))


def _style_run(run, text, pt, color, spc=None):
    run.text = text
    f = run.font
    f.name = FONT
    f.size = Pt(pt)
    f.color.rgb = RGB(color)
    if spc is not None:
        set_spc(run, spc)


def txt(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        para.line_spacing = p.get("ls", 1.15)
        if p.get("sa") is not None:
            para.space_after = Pt(p["sa"])
        if p.get("sb") is not None:
            para.space_before = Pt(p["sb"])
        for r in p["runs"]:
            run = para.add_run()
            _style_run(run, r[0], r[1], r[2], r[3] if len(r) > 3 else None)
    return box


def _no_shadow(shp):
    shp.shadow.inherit = False


def set_fill_alpha(shp, opacity_pct):
    spPr = shp._element.spPr
    sf = spPr.find(qn("a:solidFill"))
    clr = sf.find(qn("a:srgbClr"))
    a = clr.makeelement(qn("a:alpha"), {"val": str(int(opacity_pct * 1000))})
    clr.append(a)


def set_line_alpha(shp, opacity_pct):
    spPr = shp._element.spPr
    ln = spPr.find(qn("a:ln"))
    sf = ln.find(qn("a:solidFill"))
    clr = sf.find(qn("a:srgbClr"))
    a = clr.makeelement(qn("a:alpha"), {"val": str(int(opacity_pct * 1000))})
    clr.append(a)


def rect(slide, x, y, w, h, fill, line=None, line_w=0.75):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, IN(x), IN(y), IN(w), IN(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGB(fill)
    if line:
        shp.line.color.rgb = RGB(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    _no_shadow(shp)
    return shp


def card(slide, x, y, w, h, fill=GLASS, line=GLASS_LINE, radius=RADIUS, line_w=1.0,
         glass=True, fill_alpha=62, line_alpha=40):
    """Glass card: translucent smoked-blue fill + luminous border."""
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(x), IN(y), IN(w), IN(h))
    try:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGB(fill)
    if glass:
        set_fill_alpha(shp, fill_alpha)
    if line:
        shp.line.color.rgb = RGB(line)
        shp.line.width = Pt(line_w)
        if glass:
            set_line_alpha(shp, line_alpha)
    else:
        shp.line.fill.background()
    _no_shadow(shp)
    return shp


def accent_card(slide, x, y, w, h, radius=RADIUS):
    """Solid electric-blue highlight card with light border and soft glow."""
    glow = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(x - 0.055), IN(y - 0.055),
                                  IN(w + 0.11), IN(h + 0.11))
    try:
        glow.adjustments[0] = min(0.5, (radius + 0.05) / min(w, h))
    except Exception:
        pass
    glow.fill.solid()
    glow.fill.fore_color.rgb = RGB(BLUE_BRIGHT)
    set_fill_alpha(glow, 20)
    glow.line.fill.background()
    _no_shadow(glow)
    shp = card(slide, x, y, w, h, fill=ULTRA, line=ULTRA_LINE, radius=radius,
               line_w=1.25, glass=False)
    return shp


def pill(slide, x, y, w, h, fill=BTN_FILL, line=BTN_BORDER, text=None,
         size=12.5, color=BODY, spc=None, pad=0.22, glass=False, line_w=1.0):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(x), IN(y), IN(w), IN(h))
    try:
        shp.adjustments[0] = 0.5
    except Exception:
        pass
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGB(fill)
        if glass:
            set_fill_alpha(shp, 55)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = RGB(line)
        shp.line.width = Pt(line_w)
        if glass:
            set_line_alpha(shp, 55)
    else:
        shp.line.fill.background()
    _no_shadow(shp)
    if text:
        tf = shp.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = IN(pad)
        tf.margin_right = IN(pad)
        tf.margin_top = 0
        tf.margin_bottom = 0
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT
        para.line_spacing = 1.0
        run = para.add_run()
        _style_run(run, text, size, color, spc)
    return shp


def dot(slide, x, y, d, color=BLUE_BRIGHT):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, IN(x), IN(y), IN(d), IN(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGB(color)
    shp.line.fill.background()
    _no_shadow(shp)
    return shp


def cell_border(cell, color=BORDER, w_emu=9525):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    elems = []
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcPr.makeelement(qn(tag), {"w": str(w_emu), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        sf = ln.makeelement(qn("a:solidFill"), {})
        c = ln.makeelement(qn("a:srgbClr"), {"val": color})
        sf.append(c)
        ln.append(sf)
        elems.append(ln)
    for el in reversed(elems):
        tcPr.insert(0, el)


def seclabel(slide, text, w, x=ML, y=0.40):
    pill(slide, x, y, w, 0.30, fill=GLASS, line=GLASS_LINE, glass=True)
    dot(slide, x + 0.12, y + 0.105, 0.09, color=BLUE_BRIGHT)
    txt(slide, x + 0.30, y, w - 0.34, 0.30,
        [{"runs": [(text, 10, COOL, 150)], "ls": 1.0}], anchor=MSO_ANCHOR.MIDDLE)


def action_title(slide, text, y=0.78):
    txt(slide, ML, y, TITLE_W, 0.92,
        [{"runs": [(text, 26, WHITE, -90)], "ls": 1.02}])


def footer(slide, n, with_label=True, total=8):
    if with_label:
        txt(slide, ML, 7.08, 5.0, 0.24,
            [{"runs": [("GOS · SPEC-DRIVEN DEVELOPMENT", 9.5, MUTED, 150)], "ls": 1.0}])
    txt(slide, 11.28, 7.08, 1.5, 0.24,
        [{"runs": [(f"{n:02d} / {total:02d}", 9.5, MUTED, 150)], "ls": 1.0, "align": PP_ALIGN.RIGHT}])
    rect(slide, ML, 7.02, TITLE_W, 0.012, DIVIDER)


prs = Presentation()
prs.slide_width = IN(13.333)
prs.slide_height = IN(7.5)
BLANK = prs.slide_layouts[6]
_require_assets()
_slide_no = 0


def new_slide():
    global _slide_no
    _slide_no += 1
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(os.path.join(ASSETS, f"bg_{_slide_no:02d}.png"), 0, 0,
                         width=prs.slide_width, height=prs.slide_height)
    return s


# ============================ SLIDE 1 — HERO COVER ============================
s = new_slide()

pill(s, 0.55, 0.40, 3.70, 0.30, fill=GLASS, line=GLASS_LINE, glass=True)
dot(s, 0.67, 0.505, 0.09, color=BLUE_BRIGHT)
txt(s, 0.85, 0.40, 3.35, 0.30, [{"runs": [("GOS · SPEC-DRIVEN DEVELOPMENT", 10, COOL, 150)], "ls": 1.0}],
    anchor=MSO_ANCHOR.MIDDLE)

txt(s, 8.20, 0.40, 4.58, 0.30,
    [{"runs": [("ТЕХНИЧЕСКАЯ ПРЕЗЕНТАЦИЯ · ОКТЯБРЬ 2026", 10, MUTED, 150)], "ls": 1.0,
      "align": PP_ALIGN.RIGHT}], anchor=MSO_ANCHOR.MIDDLE)

txt(s, 0.85, 2.95, 8.60, 0.68, [{"runs": [("От промпта", 46, WHITE, -184)], "ls": 1.0}])
txt(s, 0.85, 3.66, 8.60, 0.68, [{"runs": [("к инженерному процессу", 46, WHITE, -184)], "ls": 1.0}])
txt(s, 0.85, 4.46, 8.60, 0.52, [{"runs": [("Spec-Driven Development с AI-агентами", 30, COOL, -90)], "ls": 1.0}])

rect(s, 0.85, 5.22, 0.04, 0.78, BLUE)
txt(s, 1.08, 5.18, 7.60, 0.88, [
    {"runs": [("Агенту можно делегировать реализацию.", 16, BODY)], "ls": 1.25, "sa": 2},
    {"runs": [("Право незаметно менять смысл задачи — нельзя.", 16, BODY)], "ls": 1.25},
])

pill(s, 10.05, 6.55, 2.73, 0.42, fill=None, line="9FC0FF", line_w=1.0)
txt(s, 10.05, 6.55, 2.73, 0.42,
    [{"runs": [("Spec → Plan → Code → Evidence", 11, WHITE)], "ls": 1.0, "align": PP_ALIGN.CENTER}],
    anchor=MSO_ANCHOR.MIDDLE)

txt(s, ML, 7.10, 1.2, 0.22, [{"runs": [("01 / 08", 9.5, MUTED, 150)], "ls": 1.0}])

# ============================ SLIDE 2 — ПРОБЛЕМА ============================
s = new_slide()
seclabel(s, "01 · ПРОБЛЕМА", 2.10)
action_title(s, "Правильно реализованная догадка всё равно остаётся догадкой")

card(s, ML, 1.80, 5.70, 3.55)
txt(s, 0.85, 2.02, 3.0, 0.25, [{"runs": [("КЕЙС", 10, COOL, 150)], "ls": 1.0}])
case_rows = [
    ("01", "Запрос: “Добавь повторный запуск неуспешной миграции”"),
    ("02", "Агент добавляет retry — все тесты проходят"),
    ("03", "После сбоя retry создаёт вторую целевую VM"),
]
for i, (num, text) in enumerate(case_rows):
    y = 2.45 + i * 0.85
    txt(s, 0.85, y, 0.62, 0.55, [{"runs": [(num, 20, BLUE_BRIGHT, -60)], "ls": 1.0}])
    txt(s, 1.52, y + 0.02, 4.50, 0.78, [{"runs": [(text, 14, BODY)], "ls": 1.15}])

txt(s, 6.55, 1.80, 5.70, 0.25, [{"runs": [("«ПОВТОР» — НЕ СОГЛАСОВАН", 10, COOL, 150)], "ls": 1.0}])
for i, t in enumerate(["Новая операция?", "Продолжение прежней?",
                       "Восстановление после частичного сбоя?", "Повторная отправка команды?"]):
    pill(s, 6.55, 2.20 + i * 0.63, 5.70, 0.50, text=t, size=13, color=BODY)

accent_card(s, ML, 5.70, TITLE_W, 0.75, radius=0.06)
txt(s, 0.90, 5.70, 11.5, 0.75,
    [{"runs": [("Ошибка возникает до написания кода:  ", 13, WHITE),
               ("idempotency · retries · partial failure · concurrency · recovery", 13, ICE)], "ls": 1.0}],
    anchor=MSO_ANCHOR.MIDDLE)

footer(s, 2)

# ============================ SLIDE 3 — ТРИ УРОВНЯ ============================
s = new_slide()
seclabel(s, "02 · МОДЕЛЬ ОШИБОК", 2.55)
action_title(s, "Ошибки возникают на трёх уровнях — контроль нужен на каждом")

levels = [
    ("УРОВЕНЬ 1 · СМЫСЛ", "Неверно поняли задачу",
     "Retry трактован как создание новой операции", "Discovery и явные решения"),
    ("УРОВЕНЬ 2 · КОНТРАКТ", "Неверно реализовали требование",
     "Дубликаты запрещены — код их создаёт", "План, TDD, spec compliance review"),
    ("УРОВЕНЬ 3 · ГОТОВНОСТЬ", "Неверно объявили готовность",
     "Happy path работает, recovery не проверен", "Evidence matrix, adversarial review"),
]
cw, gap = 3.91, 0.25
for i, (tag, head, example, ctrl) in enumerate(levels):
    x = ML + i * (cw + gap)
    card(s, x, 1.80, cw, 4.15)
    txt(s, x + 0.30, 2.06, cw - 0.60, 0.25, [{"runs": [(tag, 10, BLUE_BRIGHT, 150)], "ls": 1.0}])
    txt(s, x + 0.30, 2.42, cw - 0.60, 0.80, [{"runs": [(head, 17, WHITE, -40)], "ls": 1.08}])
    txt(s, x + 0.30, 3.30, cw - 0.60, 0.95, [{"runs": [(example, 12.5, COOL)], "ls": 1.2}])
    rect(s, x + 0.30, 4.42, cw - 0.60, 0.01, DIVIDER)
    txt(s, x + 0.30, 4.56, cw - 0.60, 0.25, [{"runs": [("КОНТРОЛЬ", 10, MUTED, 150)], "ls": 1.0}])
    txt(s, x + 0.30, 4.84, cw - 0.60, 0.80, [{"runs": [(ctrl, 12.5, SKY)], "ls": 1.2}])

txt(s, ML, 6.25, TITLE_W, 0.35,
    [{"runs": [("Одного code review недостаточно: смысл, реализация и готовность проверяются отдельно",
                13, COOL)], "ls": 1.0}])
footer(s, 3)

# ============================ SLIDE 4 — КОНТРАКТ ============================
s = new_slide()
seclabel(s, "03 · КОНТРАКТ", 2.00)
action_title(s, "Спецификация определяет обязательства, план организует работу, код исполняет")

chain = [
    ("Proposal", "зачем изменение и его границы"),
    ("Specs", "обязательное поведение и гарантии"),
    ("Design", "как реализовать, какие trade-offs"),
    ("Tasks", "работа внутри change"),
    ("Implementation plan", "файлы, тесты, шаги исполнителя"),
]
rect(s, 0.715, 2.10, 0.016, 3.42, BLUE)
for i, (name, role) in enumerate(chain):
    y = 2.00 + i * 0.78
    dot(s, 0.665, y + 0.10, 0.115, color=BLUE_BRIGHT)
    txt(s, 0.98, y, 5.10, 0.30, [{"runs": [(name, 14, WHITE, -40)], "ls": 1.0}])
    txt(s, 0.98, y + 0.30, 5.10, 0.28, [{"runs": [(role, 11.5, COOL)], "ls": 1.0}])
txt(s, 0.98, 5.92, 5.40, 0.26, [{"runs": [("План не становится вторым источником требований",
                                           10, MUTED)], "ls": 1.0}])

card(s, 7.10, 1.80, 5.13, 4.30)
txt(s, 7.40, 2.02, 4.55, 0.25, [{"runs": [("ХОРОШЕЕ ТРЕБОВАНИЕ", 10, COOL, 150)], "ls": 1.0}])
txt(s, 7.40, 2.42, 4.55, 0.50, [{"runs": [("✗   ", 12.5, MUTED),
                                          ("Повторный запуск должен работать корректно", 12.5, MUTED)],
                                 "ls": 1.2}])
txt(s, 7.40, 3.02, 4.55, 1.15, [{"runs": [("✓   ", 12.5, BLUE_BRIGHT),
                                          ("Если операция с данным ID уже создала ресурс, повтор "
                                           "продолжает её или возвращает результат — но не создаёт "
                                           "новый ресурс", 12.5, BODY)], "ls": 1.2}])
rect(s, 7.40, 4.42, 4.55, 0.01, DIVIDER)
txt(s, 7.40, 4.58, 4.55, 0.80,
    [{"runs": [("Открываются проверяемые вопросы: идентичность операции · конкуренция · "
                "потерянный ответ API", 10, MUTED)], "ls": 1.25}])
footer(s, 4)

# ============================ SLIDE 5 — FACTS + SPEC_DRIFT ============================
s = new_slide()
seclabel(s, "04 · DISCOVERY", 2.15)
action_title(s, "Факты ≠ решения. Менять контракт можно, незаметно — нельзя")

facts = [
    ("FACT", "У операции есть устойчивый operation ID"),
    ("INFERENCE", "Поле нужно для восстановления — вывод, не контракт"),
    ("UNKNOWN", "Создался ли внешний ресурс до timeout"),
    ("CONFLICT", "Документация обещает resume, код допускает restart"),
    ("DECISION", "Повтор использует исходный operation ID"),
    ("ASSUMPTION", "Внешний API находит ранее созданный ресурс"),
]
fw, fh, fgx, fgy = 2.98, 1.42, 0.22, 0.22
for i, (tag, text) in enumerate(facts):
    x = ML + (i % 2) * (fw + fgx)
    y = 1.80 + (i // 2) * (fh + fgy)
    card(s, x, y, fw, fh)
    txt(s, x + 0.22, y + 0.16, fw - 0.44, 0.25, [{"runs": [(tag, 11, BLUE_BRIGHT, 150)], "ls": 1.0}])
    txt(s, x + 0.22, y + 0.46, fw - 0.44, 0.85, [{"runs": [(text, 10.5, COOL)], "ls": 1.2}])

txt(s, 7.10, 1.80, 5.13, 0.25, [{"runs": [("SPEC_DRIFT", 10, COOL, 150)], "ls": 1.0}])
drift = ["Противоречие при реализации", "Evidence + варианты + последствия", "Явное решение",
         "Обновление specs / design / tasks", "Validation → продолжение"]
rect(s, 7.31, 2.46, 0.015, 2.80, BLUE)
for i, t in enumerate(drift):
    y = 2.20 + i * 0.70
    dot(s, 7.26, y + 0.20, 0.10, color=BLUE_BRIGHT)
    pill(s, 7.10, y, 5.13, 0.52, text=t, size=11.5, color=BODY, pad=0.40)
accent_card(s, 7.10, 5.72, 5.13, 0.85, radius=0.06)
txt(s, 7.40, 5.72, 4.60, 0.85,
    [{"runs": [("Freeze — не запрет изменить решение. Это запрет изменить его незаметно",
                12.5, WHITE)], "ls": 1.15}], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 5)

# ============================ SLIDE 6 — ИСПОЛНЕНИЕ ============================
s = new_slide()
seclabel(s, "05 · ИСПОЛНЕНИЕ", 2.35)
action_title(s, "Исполнение проверяется дважды: соответствие контракту и качество")

roles = [
    ("Координатор", "карта задач и зависимостей · передача контракта · организация review"),
    ("Исполнитель", "ограниченная задача · запускает проверки · отчёт о результате"),
    ("Reviewer", "соответствие контракту · качество реализации"),
]
for i, (role, duty) in enumerate(roles):
    y = 1.80 + i * 1.42
    card(s, ML, y, 6.20, 1.30)
    txt(s, 0.85, y + 0.18, 5.6, 0.32, [{"runs": [(role, 14, WHITE, -40)], "ls": 1.0}])
    txt(s, 0.85, y + 0.56, 5.60, 0.62, [{"runs": [(duty, 11.5, COOL)], "ls": 1.2}])
txt(s, ML, 6.14, 6.20, 0.30,
    [{"runs": [("Субагент — граница контекста и ответственности, а не «виртуальная команда»",
                10, MUTED)], "ls": 1.0}])

card(s, 7.10, 1.80, 5.13, 4.20)
pill(s, 7.35, 2.10, 1.10, 0.44, fill=ULTRA, line=ULTRA_LINE, text="RED", size=12.5,
     color=WHITE, pad=0.18, line_w=1.0)
txt(s, 8.49, 2.10, 0.22, 0.44, [{"runs": [("→", 12, MUTED)], "ls": 1.0, "align": PP_ALIGN.CENTER}],
    anchor=MSO_ANCHOR.MIDDLE)
pill(s, 8.75, 2.10, 1.40, 0.44, text="GREEN", size=12.5, color=BODY, pad=0.18)
txt(s, 10.19, 2.10, 0.22, 0.44, [{"runs": [("→", 12, MUTED)], "ls": 1.0, "align": PP_ALIGN.CENTER}],
    anchor=MSO_ANCHOR.MIDDLE)
pill(s, 10.45, 2.10, 1.53, 0.44, text="REFACTOR", size=12.5, color=BODY, pad=0.18)

txt(s, 7.35, 2.95, 4.55, 0.25, [{"runs": [("ПОЛЕЗНЫЙ RED", 10, COOL, 150)], "ls": 1.0}])
txt(s, 7.35, 3.22, 4.55, 0.45, [{"runs": [("Обнаружены два целевых ресурса вместо одного",
                                           12.5, BODY)], "ls": 1.15}])
rect(s, 7.35, 3.76, 4.55, 0.01, DIVIDER)
txt(s, 7.35, 3.92, 4.55, 0.25, [{"runs": [("БЕСПОЛЕЗНЫЙ RED", 10, MUTED, 150)], "ls": 1.0}])
txt(s, 7.35, 4.19, 4.55, 0.45, [{"runs": [("ImportError в самом тесте", 12.5, COOL)], "ls": 1.15}])
footer(s, 6)

# ============================ SLIDE 7 — EVIDENCE ============================
s = new_slide()
seclabel(s, "06 · ДОКАЗАТЕЛЬСТВА", 2.80)
action_title(s, "Готовность объявляется только по свежим доказательствам")

rows = [
    ("R-01 · повтор сохраняет operation ID", "unit / contract test"),
    ("R-02 · повтор не создаёт дубликат", "integration fault test"),
    ("R-03 · tenant boundary сохраняется", "negative integration test"),
    ("R-04 · неоднозначность видна явно", "failure-path test"),
]
gf = s.shapes.add_table(5, 3, IN(ML), IN(1.80), IN(7.20), IN(3.30))
table = gf.table
table.first_row = False
table.horz_banding = False
table.columns[0].width = IN(3.60)
table.columns[1].width = IN(2.20)
table.columns[2].width = IN(1.40)
table.rows[0].height = IN(0.42)
for i in range(1, 5):
    table.rows[i].height = IN(0.72)

headers = ["ТРЕБОВАНИЕ", "ПРОВЕРКА", "СТАТУС"]
for c, htext in enumerate(headers):
    cell = table.cell(0, c)
    cell_border(cell, color="3A5182")
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGB(BG3)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = IN(0.16)
    cell.margin_right = IN(0.10)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    _style_run(r, htext, 10, COOL, 150)

for ri, (req, ver) in enumerate(rows, start=1):
    for ci, val in enumerate([req, ver, "NOT RUN"]):
        cell = table.cell(ri, ci)
        cell_border(cell)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGB(CARD_ALT)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = IN(0.16)
        cell.margin_right = IN(0.10)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        color = MUTED if ci == 2 else BODY
        r = p.add_run()
        _style_run(r, val, 11.5, color)

rect(s, ML, 2.94, 0.05, 0.72, ULTRA)
txt(s, ML, 5.28, 7.20, 0.26,
    [{"runs": [("Статусы иллюстративные: проверки не выполнялись", 10, MUTED)], "ls": 1.0}])

card(s, 8.05, 1.80, 4.18, 3.30)
finding = [
    ("CLAIM", "Повтор не создаёт вторую целевую VM", 2.00, 0.50),
    ("EVIDENCE REQUIRED", "Сбой после внешнего create, но до сохранения", 2.86, 0.50),
    ("EVIDENCE FOUND", "Только тест повтора после успешного commit", 3.72, 0.50),
]
for label, val, y, hh in finding:
    txt(s, 8.35, y, 3.60, 0.25, [{"runs": [(label, 10, BLUE_BRIGHT, 150)], "ls": 1.0}])
    txt(s, 8.35, y + 0.25, 3.60, hh, [{"runs": [(val, 11.5, BODY)], "ls": 1.15}])
accent_card(s, 8.35, 4.46, 1.10, 0.44, radius=0.15)
txt(s, 8.35, 4.46, 1.10, 0.44, [{"runs": [("GAP", 14, WHITE, 150)], "ls": 1.0,
                                 "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)

txt(s, ML, 6.10, TITLE_W, 0.35,
    [{"runs": [("Тесты могут быть зелёными, а требование — всё ещё непроверенным", 13, COOL)], "ls": 1.0}])
footer(s, 7)

# ============================ SLIDE 8 — ВНЕДРЕНИЕ ============================
s = new_slide()
seclabel(s, "07 · ВНЕДРЕНИЕ", 2.15)
action_title(s, "Масштабировать следует церемонию, а не ответственность")

phases = [
    ("ПОНЯТЬ", "Система и evidence"),
    ("СОГЛАСОВАТЬ", "Discovery и решения"),
    ("ЗАФИКСИРОВАТЬ", "Binding specification"),
    ("РЕАЛИЗОВАТЬ", "План, TDD, субагенты"),
    ("ПОДТВЕРДИТЬ", "Fresh verification, review"),
    ("ИНТЕГРИРОВАТЬ", "Archive, merge, PR"),
]
pw = (TITLE_W - 5 * 0.12) / 6
for i, (label, card_text) in enumerate(phases):
    x = ML + i * (pw + 0.12)
    active = i == 2
    txt(s, x, 1.72, pw, 0.22,
        [{"runs": [(label, 10, WHITE if active else MUTED, 150)], "ls": 1.0, "align": PP_ALIGN.CENTER}])
    if active:
        glow = card(s, x - 0.06, 1.96, pw + 0.12, 1.62, fill=BLUE_BRIGHT, line=None,
                    radius=0.16, glass=True, fill_alpha=18)
    c = card(s, x, 2.02, pw, 1.50, fill=ULTRA if active else GLASS,
             line=ULTRA_LINE if active else DIVIDER, line_w=1.25 if active else 1.0,
             glass=not active, fill_alpha=88 if not active else 62)
    txt(s, x + 0.16, 2.20, pw - 0.32, 1.15,
        [{"runs": [(card_text, 11.5, WHITE if active else BODY)], "ls": 1.2,
          "align": PP_ALIGN.CENTER if active else PP_ALIGN.LEFT}])

modes = [
    ("LIGHTWEIGHT", "локальные обратимые изменения"),
    ("STANDARD", "API, бизнес-логика, интеграции"),
    ("HIGH-RISK", "данные, security, tenancy"),
    ("INCIDENT", "стабилизация, rollback, затем reconciliation"),
]
mw, mh = 3.49, 1.15
for i, (tag, text) in enumerate(modes):
    x = ML + (i % 2) * (mw + 0.22)
    y = 4.30 + (i // 2) * (mh + 0.17)
    card(s, x, y, mw, mh)
    txt(s, x + 0.24, y + 0.14, mw - 0.48, 0.25, [{"runs": [(tag, 11, BLUE_BRIGHT, 150)], "ls": 1.0}])
    txt(s, x + 0.24, y + 0.44, mw - 0.48, 0.60, [{"runs": [(text, 10.5, COOL)], "ls": 1.2}])

accent_card(s, 8.05, 4.30, 4.18, 1.30, radius=0.06)
txt(s, 8.29, 4.44, 3.70, 0.66,
    [{"runs": [("Главный результат процесса — не больше документов, а меньше скрытых решений",
                13, WHITE)], "ls": 1.15}])
txt(s, 8.29, 5.14, 3.70, 0.36,
    [{"runs": [("Старт: один пилот · закреплённые версии · честные статусы GAP", 10, ICE)], "ls": 1.15}])
footer(s, 8)

# ============================ save + verify ============================
prs.core_properties.title = "От промпта к инженерному процессу: Spec-Driven Development с AI-агентами"
prs.core_properties.subject = "Техническая презентация, октябрь 2026"
prs.save(OUT_PPTX)
print(f"saved {OUT_PPTX}, slides: {len(prs.slides._sldIdLst)}")


def verify():
    issues = []
    sw, sh = prs.slide_width, prs.slide_height
    for idx, slide in enumerate(prs.slides, start=1):
        boxes = []
        for shp in slide.shapes:
            if shp.left is None:
                continue
            l, t, w, h = shp.left, shp.top, shp.width, shp.height
            if l < -1000 or t < -1000 or l + w > sw + 1000 or t + h > sh + 1000:
                issues.append(f"slide {idx}: shape out of bounds: {shp.shape_type} "
                              f"({l / EMU_IN:.2f},{t / EMU_IN:.2f},{w / EMU_IN:.2f},{h / EMU_IN:.2f})")
            has_text = shp.has_text_frame and shp.text_frame.text.strip()
            if has_text and shp.shape_type is not None:
                boxes.append((l, t, w, h, shp.text_frame.text[:28]))
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                ix = max(0, min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0]))
                iy = max(0, min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]))
                inter = (ix / EMU_IN) * (iy / EMU_IN)
                if inter > 0.05:
                    issues.append(f"slide {idx}: text overlap {inter:.2f} sq in: "
                                  f"«{a[4]}» × «{b[4]}»")
    if issues:
        print("VERIFY ISSUES:")
        for i in issues:
            print(" -", i)
    else:
        print("verify OK: bounds + text overlaps clean")


verify()
