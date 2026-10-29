"""pptx-ai-slide-forge :: text QA + geometry lint for a built deck.

Usage:
    python3 qa_check.py deck.pptx --slides 8 \
        --contains "NOT RUN:4" --contains "GOT IT" --forbid "Спасибо"

Checks:
  1. Slide count equals --slides.
  2. Every --contains string present (optionally "TEXT:COUNT" for an exact count).
  3. Every --forbid string absent.
  4. Geometry lint: no shape outside slide bounds; no two text-bearing shapes
     overlapping by more than --max-overlap square inches.

Text extraction uses markitdown when available and falls back to python-pptx.
Exits non-zero if any check fails. All numbers reported are from real runs —
never report a pass you have not executed.
"""
import argparse
import os
import sys

from pptx import Presentation

EMU_IN = 914400


def text_dump(path: str) -> str:
    if not os.path.isfile(path):
        raise SystemExit(f"deck not found: {path}")
    try:
        from markitdown import MarkItDown
        return MarkItDown().convert(path).text_content
    except Exception as e:
        print(f"[warn] markitdown failed ({e}); falling back to python-pptx text walk")
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"<!-- Slide number: {i} -->")
        for shp in slide.shapes:
            if shp.has_text_frame:
                parts.append(shp.text_frame.text)
    return "\n".join(parts)


def geometry_lint(path: str, max_overlap: float):
    prs = Presentation(path)
    issues = []
    sw, sh = prs.slide_width, prs.slide_height
    for idx, slide in enumerate(prs.slides, start=1):
        boxes = []
        for shp in slide.shapes:
            if shp.left is None:
                continue
            l, t, w, h = shp.left, shp.top, shp.width, shp.height
            if l < -1000 or t < -1000 or l + w > sw + 1000 or t + h > sh + 1000:
                issues.append(
                    f"slide {idx}: shape out of bounds "
                    f"({l / EMU_IN:.2f},{t / EMU_IN:.2f},{w / EMU_IN:.2f},{h / EMU_IN:.2f})")
            if shp.has_text_frame and shp.text_frame.text.strip():
                boxes.append((l, t, w, h, shp.text_frame.text[:28]))
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                ix = max(0, min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0]))
                iy = max(0, min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]))
                inter = (ix / EMU_IN) * (iy / EMU_IN)
                if inter > max_overlap:
                    issues.append(f"slide {idx}: text overlap {inter:.2f} sq in: "
                                  f"\u00ab{a[4]}\u00bb \u00d7 \u00ab{b[4]}\u00bb")
    return issues, len(prs.slides._sldIdLst)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("deck", help="path to .pptx")
    ap.add_argument("--slides", type=int, default=None, help="expected slide count")
    ap.add_argument("--contains", action="append", default=[],
                    help='required substring, optionally "TEXT:COUNT" for exact count')
    ap.add_argument("--forbid", action="append", default=[], help="forbidden substring")
    ap.add_argument("--max-overlap", type=float, default=0.05,
                    help="text overlap threshold, square inches (default 0.05)")
    args = ap.parse_args()

    failures = []

    text = text_dump(args.deck)

    if args.slides is not None:
        import re
        n = len(re.findall(r"<!-- Slide number:", text)) or \
            text.count("<!-- Slide")
        ok = n == args.slides
        print(f"[{'OK' if ok else 'FAIL'}] slide count: {n} (expected {args.slides})")
        if not ok:
            failures.append("slide count")

    for spec in args.contains:
        if ":" in spec and spec.rsplit(":", 1)[1].isdigit():
            needle, want = spec.rsplit(":", 1)
            want = int(want)
        else:
            needle, want = spec, None
        got = text.count(needle)
        ok = (got == want) if want is not None else (got > 0)
        print(f"[{'OK' if ok else 'FAIL'}] contains {needle!r}: {got}"
              + (f" (expected {want})" if want is not None else ""))
        if not ok:
            failures.append(f"contains {needle!r}")

    for needle in args.forbid:
        got = text.count(needle)
        ok = got == 0
        print(f"[{'OK' if ok else 'FAIL'}] forbid {needle!r}: {got}")
        if not ok:
            failures.append(f"forbid {needle!r}")

    issues, n_slides = geometry_lint(args.deck, args.max_overlap)
    if issues:
        print(f"[FAIL] geometry lint ({n_slides} slides):")
        for i in issues:
            print("  -", i)
        failures.append("geometry lint")
    else:
        print(f"[OK] geometry lint ({n_slides} slides): bounds + overlaps clean")

    if failures:
        print(f"\nQA FAIL: {len(failures)} failing checks")
        sys.exit(1)
    print("\nQA PASS: all checks green")


if __name__ == "__main__":
    main()
