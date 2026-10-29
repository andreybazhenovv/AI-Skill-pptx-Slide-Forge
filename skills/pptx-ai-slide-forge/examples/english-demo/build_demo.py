"""pptx-ai-slide-forge :: compact 3-slide English demo deck.

A minimal, self-contained build showing the skill's visual language on
English copy: liquid-ribbon backgrounds + glass cards + electric accents.
The full reference implementation (helpers, footers, tables, geometry lint)
lives in scripts/build_deck.py.

Run from this directory:
    python3 build_demo.py
Then render for visual QA:
    bash ../../scripts/render_qa.sh "$PWD/english-demo.pptx" qa
"""
import os

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.abspath(os.path.join(HERE, "..", "..", "assets"))
OUT = os.path.join(HERE, "english-demo.pptx")
EMU_IN = 914400


def IN(v):
    return Emu(int(round(v * EMU_IN)))


# ---- NovaPay Liquid Electric Blue tokens ----
BG2, GLASS, GLASS_LINE = "0A1830", "20365B", "8AB4FF"
BLUE, BLUE_BRIGHT, ULTRA, ULTRA_LINE = "075BFF", "287BFF", "2438F5", "C6D6FF"
ICE, WHITE, COOL, MUTED, BODY = "B5D7FF", "F8FAFF", "A5B7D9", "6F82A8", "E2E9FA"
DIVIDER = "22335A"
FONT = "Helvetica Neue"
RADIUS = 0.125
ML, TITLE_W = 0.55, 12.233

prs = Presentation()
prs.slide_width, prs.slide_height = IN(13.333), IN(7.5)
BLANK = prs.slide_layouts[6]
_no = 0


def RGB(h):
    return RGBColor.from_string(h)


def set_spc(run, spc):
    run._r.get_or_add_rPr().set("spc", str(int(spc)))


def run_style(run, text, pt, color, spc=None):
    run.text = text
    run.font.name, run.font.size = FONT, Pt(pt)
    run.font.color.rgb = RGB(color)
    if spc is not None:
        set_spc(run, spc)


def txt(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        para.line_spacing = p.get("ls", 1.15)
        for r in p["runs"]:
            run_style(para.add_run(), r[0], r[1], r[2], r[3] if len(r) > 3 else None)
    return box


def no_shadow(shp):
    shp.shadow.inherit = False


def set_fill_alpha(shp, pct):
    sf = shp._element.spPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")
    clr = sf.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    clr.append(clr.makeelement(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha",
        {"val": str(int(pct * 1000))}))


def set_line_alpha(shp, pct):
    ln = shp._element.spPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}ln")
    sf = ln.find("{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")
    clr = sf.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    clr.append(clr.makeelement(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha",
        {"val": str(int(pct * 1000))}))


def rect(slide, x, y, w, h, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, IN(x), IN(y), IN(w), IN(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGB(color)
    shp.line.fill.background()
    no_shadow(shp)
    return shp


def card(slide, x, y, w, h, fill=GLASS, line=GLASS_LINE, radius=RADIUS, line_w=1.0,
         glass=True, fill_alpha=62):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(x), IN(y), IN(w), IN(h))
    try:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGB(fill)
    if glass:
        set_fill_alpha(shp, fill_alpha)
    if line:
        shp.line.color.rgb = RGB(line)
        shp.line.width = Pt(line_w)
        if glass:
            set_line_alpha(shp, 40)
    else:
        shp.line.fill.background()
    no_shadow(shp)
    return shp


def accent_card(slide, x, y, w, h, radius=RADIUS):
    glow = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(x - 0.055), IN(y - 0.055),
                                  IN(w + 0.11), IN(h + 0.11))
    try:
        glow.adjustments[0] = min(0.5, (radius + 0.05) / min(w, h))
    except Exception:
        pass
    glow.fill.solid()
    glow.fill.fore_color.rgb = RGB(BLUE_BRIGHT)
    set_fill_alpha(glow, 20)
    glow.line.fill.background()
    no_shadow(glow)
    return card(slide, x, y, w, h, fill=ULTRA, line=ULTRA_LINE, radius=radius,
                line_w=1.25, glass=False)


def pill(slide, x, y, w, h, fill=GLASS, line=GLASS_LINE, text=None, size=12.5,
         color=BODY, pad=0.22, glass=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, IN(x), IN(y), IN(w), IN(h))
    try:
        shp.adjustments[0] = 0.5
    except Exception:
        pass
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGB(fill)
        if glass:
            set_fill_alpha(shp, 55)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = RGB(line)
        shp.line.width = Pt(1.0)
        if glass:
            set_line_alpha(shp, 55)
    else:
        shp.line.fill.background()
    no_shadow(shp)
    if text:
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left, tf.margin_right = IN(pad), IN(pad)
        tf.margin_top = tf.margin_bottom = 0
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT
        para.line_spacing = 1.0
        run_style(para.add_run(), text, size, color)
    return shp


def dot(slide, x, y, d=0.09, color=BLUE_BRIGHT):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, IN(x), IN(y), IN(d), IN(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGB(color)
    shp.line.fill.background()
    no_shadow(shp)


def new_slide(bg):
    global _no
    _no += 1
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(os.path.join(ASSETS, bg), 0, 0,
                         width=prs.slide_width, height=prs.slide_height)
    return s


def chip(slide, text, w, x=ML, y=0.40):
    pill(slide, x, y, w, 0.30, text=text, size=10, color=COOL, pad=0.30)
    dot(slide, x + 0.12, y + 0.105)


def title(slide, text):
    txt(slide, ML, 0.78, TITLE_W, 0.92, [{"runs": [(text, 26, WHITE, -90)], "ls": 1.02}])


def footer(slide, n):
    txt(slide, ML, 7.08, 5.0, 0.24,
        [{"runs": [("GOS · SPEC-DRIVEN DEVELOPMENT", 9.5, MUTED, 150)], "ls": 1.0}])
    txt(slide, 11.28, 7.08, 1.5, 0.24,
        [{"runs": [(f"{n:02d} / 03", 9.5, MUTED, 150)], "ls": 1.0, "align": PP_ALIGN.RIGHT}])
    rect(slide, ML, 7.02, TITLE_W, 0.012, DIVIDER)


# ============================ SLIDE 1 — HERO COVER ============================
s = new_slide("bg_01.png")
pill(s, 0.55, 0.40, 3.70, 0.30, text="GOS · SPEC-DRIVEN DEVELOPMENT", size=10, color=COOL, pad=0.30)
dot(s, 0.67, 0.505)
txt(s, 8.20, 0.40, 4.58, 0.30,
    [{"runs": [("TECHNICAL TALK · OCTOBER 2026", 10, MUTED, 150)], "ls": 1.0,
      "align": PP_ALIGN.RIGHT}], anchor=MSO_ANCHOR.MIDDLE)
txt(s, 0.85, 2.95, 8.60, 0.68, [{"runs": [("From prompts", 46, WHITE, -184)], "ls": 1.0}])
txt(s, 0.85, 3.66, 8.60, 0.68, [{"runs": [("to an engineering process", 46, WHITE, -184)], "ls": 1.0}])
txt(s, 0.85, 4.46, 8.60, 0.52,
    [{"runs": [("Spec-Driven Development with AI agents", 30, COOL, -90)], "ls": 1.0}])
rect(s, 0.85, 5.22, 0.04, 0.78, BLUE)
txt(s, 1.08, 5.18, 7.60, 0.88, [
    {"runs": [("Implementation can be delegated to an agent.", 16, BODY)], "ls": 1.25, "sa": 2},
    {"runs": [("The right to silently change the meaning of a task — cannot.", 16, BODY)], "ls": 1.25},
])
pill(s, 10.05, 6.55, 2.73, 0.42, fill=None, line="9FC0FF")
txt(s, 10.05, 6.55, 2.73, 0.42,
    [{"runs": [("Spec → Plan → Code → Evidence", 11, WHITE)], "ls": 1.0, "align": PP_ALIGN.CENTER}],
    anchor=MSO_ANCHOR.MIDDLE)
txt(s, ML, 7.10, 1.2, 0.22, [{"runs": [("01 / 03", 9.5, MUTED, 150)], "ls": 1.0}])

# ============================ SLIDE 2 — THE PROBLEM ============================
s = new_slide("bg_02.png")
chip(s, "01 · THE PROBLEM", 2.10)
title(s, "A correctly implemented guess is still a guess")

card(s, ML, 1.80, 5.70, 3.55)
txt(s, 0.85, 2.02, 3.0, 0.25, [{"runs": [("THE CASE", 10, COOL, 150)], "ls": 1.0}])
for i, (num, text) in enumerate([
        ("01", "Request: \u201cAdd a retry for failed migrations\u201d"),
        ("02", "The agent adds a retry — all tests pass"),
        ("03", "After a failure the retry creates a second target VM")]):
    y = 2.45 + i * 0.85
    txt(s, 0.85, y, 0.62, 0.55, [{"runs": [(num, 20, BLUE_BRIGHT, -60)], "ls": 1.0}])
    txt(s, 1.52, y + 0.02, 4.50, 0.78, [{"runs": [(text, 14, BODY)], "ls": 1.15}])

txt(s, 6.55, 1.80, 5.70, 0.25,
    [{"runs": [("\u201cRETRY\u201d — NOT AGREED", 10, COOL, 150)], "ls": 1.0}])
for i, t in enumerate(["A new operation?", "Continuation of the previous one?",
                       "Recovery after a partial failure?", "Resending the command?"]):
    pill(s, 6.55, 2.20 + i * 0.63, 5.70, 0.50, text=t, size=13, color=BODY)

accent_card(s, ML, 5.70, TITLE_W, 0.75, radius=0.06)
txt(s, 0.90, 5.70, 11.5, 0.75,
    [{"runs": [("The error happens before code is written:  ", 13, WHITE),
               ("idempotency · retries · partial failure · concurrency · recovery", 13, ICE)],
      "ls": 1.0}], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2)

# ============================ SLIDE 3 — THREE LEVELS ============================
s = new_slide("bg_03.png")
chip(s, "02 · ERROR MODEL", 2.30)
title(s, "Errors occur at three levels — each level needs its own control")

levels = [
    ("LEVEL 1 · MEANING", "The task was misunderstood",
     "Retry treated as creating a new operation", "Discovery and explicit decisions"),
    ("LEVEL 2 · CONTRACT", "The requirement was misimplemented",
     "Duplicates are forbidden — the code creates them", "Plan, TDD, spec compliance review"),
    ("LEVEL 3 · READINESS", "Readiness was declared wrongly",
     "Happy path works, recovery is untested", "Evidence matrix, adversarial review"),
]
cw, gap = 3.91, 0.25
for i, (tag, head, example, ctrl) in enumerate(levels):
    x = ML + i * (cw + gap)
    card(s, x, 1.80, cw, 4.15)
    txt(s, x + 0.30, 2.06, cw - 0.60, 0.25, [{"runs": [(tag, 10, BLUE_BRIGHT, 150)], "ls": 1.0}])
    txt(s, x + 0.30, 2.42, cw - 0.60, 0.80, [{"runs": [(head, 17, WHITE, -40)], "ls": 1.08}])
    txt(s, x + 0.30, 3.30, cw - 0.60, 0.95, [{"runs": [(example, 12.5, COOL)], "ls": 1.2}])
    rect(s, x + 0.30, 4.42, cw - 0.60, 0.01, DIVIDER)
    txt(s, x + 0.30, 4.56, cw - 0.60, 0.25, [{"runs": [("CONTROL", 10, MUTED, 150)], "ls": 1.0}])
    txt(s, x + 0.30, 4.84, cw - 0.60, 0.80, [{"runs": [(ctrl, 12.5, SKY := "5BA7FF")], "ls": 1.2}])

txt(s, ML, 6.25, TITLE_W, 0.35,
    [{"runs": [("One code review is not enough: meaning, implementation and readiness "
                "are verified separately", 13, COOL)], "ls": 1.0}])
footer(s, 3)

prs.core_properties.title = "From prompts to an engineering process (demo)"
prs.save(OUT)
print(f"saved {OUT}, slides: {len(prs.slides._sldIdLst)}")

# quick geometry sanity: bounds + text overlaps
issues, sw, sh = [], prs.slide_width, prs.slide_height
for idx, slide in enumerate(prs.slides, start=1):
    boxes = []
    for shp in slide.shapes:
        if shp.left is None:
            continue
        if shp.left < -1000 or shp.top < -1000 or shp.left + shp.width > sw + 1000 \
                or shp.top + shp.height > sh + 1000:
            issues.append(f"slide {idx}: out of bounds")
        if shp.has_text_frame and shp.text_frame.text.strip():
            boxes.append((shp.left, shp.top, shp.width, shp.height, shp.text_frame.text[:24]))
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ix = max(0, min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0]))
            iy = max(0, min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]))
            if (ix / EMU_IN) * (iy / EMU_IN) > 0.05:
                issues.append(f"slide {idx}: overlap \u00ab{a[4]}\u00bb \u00d7 \u00ab{b[4]}\u00bb")
print("geometry:", issues if issues else "clean")
